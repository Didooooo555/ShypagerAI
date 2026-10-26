from flask import Flask, render_template, request, redirect, url_for, session, send_file
from openai import OpenAI
from docx import Document
from openpyxl import Workbook
import mysql.connector
import bcrypt
import base64
import os

app = Flask(__name__)
app.secret_key = ''

client = OpenAI(api_key="")
MODEL = "gpt-4"

@app.route('/')
def index():
    return render_template('glav.html')

@app.route('/regis', methods=['GET', 'POST'])
def regis():
    return render_template('regis.html', request=request)

@app.route('/submit', methods=['POST'])
def submit():
    if request.method == 'POST':
        con = mysql.connector.connect(
            host='localhost',
            user='root',
            password='didar09k',
            database='ty'
        )
        cursor = con.cursor()
 
       
        cursor.execute('''CREATE TABLE IF NOT EXISTS user (
            id INT AUTO_INCREMENT PRIMARY KEY,
            username VARCHAR(100),
            surname VARCHAR(100),
            age INT,
            rost INT,
            ves INT,
            pol VARCHAR(100)
        )''')

        cursor.execute('''CREATE TABLE IF NOT EXISTS user_iin (
            id INT AUTO_INCREMENT PRIMARY KEY,
            iin VARBINARY(60) UNIQUE
        )''')


        username = request.form['username']
        surname = request.form['surname']
        age = request.form['age']
        rost = request.form['rost']
        ves = request.form['ves']
        pol = request.form['pol']
        iin = request.form['iin']

        session['username'] = username
        session['age'] = age
        session['rost'] = rost
        session['ves'] = ves
        session['pol'] = pol
        session['iin'] = iin

        cursor.execute('SELECT iin FROM user_iin')
        existing_hashes = cursor.fetchall()

        exists = False
        for hash_bytes in existing_hashes:
            if bcrypt.checkpw(iin.encode('utf-8'), hash_bytes[0]):  
                exists = True
                break

        if exists:
            return redirect(url_for('osnov'))
        else:
            hashed_iin = bcrypt.hashpw(iin.encode('utf-8'), bcrypt.gensalt())
            cursor.execute('INSERT INTO user_iin (iin) VALUES (%s)', (hashed_iin,))
            cursor.execute('INSERT INTO user (username,surname,age,rost,ves,pol) VALUES(%s,%s,%s,%s,%s,%s)', (username,surname,age,rost,ves,pol))
            con.commit()
            
            

            cursor.execute("SELECT * FROM user")
            users_data = cursor.fetchall()
            column_names = [desc[0] for desc in cursor.description]

            wb = Workbook()
            ws = wb.active
            ws.title = "Users"
            ws.append(column_names)
            for row in users_data:
                ws.append(row)

            temp_folder = os.path.join(os.getcwd(), "temp")
            if not os.path.exists(temp_folder):
                os.makedirs(temp_folder)

            safe_username = "".join([c if c.isalnum() else "_" for c in username])
            safe_surname = "".join([c if c.isalnum() else "_" for c in surname])
            excel_filename = f"user_{safe_username}_{safe_surname}.xlsx"
            excel_filepath = os.path.join(temp_folder, excel_filename)

            wb.save(excel_filepath)
            session['excel_path'] = excel_filepath

        cursor.close()
        con.close()
    return redirect(url_for('osnov'))

@app.route('/osnov', methods=['GET', 'POST'])
def osnov():
    result = None
    data_points = session.get('data_points', [])

    age = session.get('age', '')
    rost = session.get('rost', '')
    ves = session.get('ves', '')
    pol = session.get('pol', '')

    if request.method == 'POST':
        symptoms = request.form['symptoms']
        duration = request.form['duration']
        intensity = request.form['intensity']
        image = request.files.get('image')

        data_points.append({
            'symptoms': symptoms,
            'duration': duration,
            'intensity': intensity
        })
        session['data_points'] = data_points

        prompt = f"""Пациент жалуется на: {symptoms}. Длительность: {duration} дней. Интенсивность: {intensity}/10. 
Возраст: {age}, вес: {ves}, рост: {rost}, пол: {pol}. 
Проанализируй симптомы и изображение (если есть). 
Дай возможный диагноз и рекомендации по лечению или возможные заболевания, легкие советы питанию, сну. Укажи, какие лекарства могут помочь (если доступны без рецепта).
Пожалуйста, учитывай весь анализ пациента. Дай рекомендаций как избежать других последствий. Можешь укоротить советы и другие, но возможные заболевания и диагноз не трогай.
После этого всего обязательно добавь что ты не врачь но коротко."""

        messages = [
            {"role": "system", "content": "Ты доброжелательный медицинский помощник. Отвечай понятно и аккуратно."},
            {"role": "user", "content": prompt}
        ]

        if image and image.filename:
            image_data = base64.b64encode(image.read()).decode("utf-8")
            image_input = {
                "type": "image_url",
                "image_url": {
                    "url": f"data:image/jpeg;base64,{image_data}",
                    "detail": "high"
                }
            }
            messages.append({"role": "user", "content": image_input})

        try:
            chat_completion = client.chat.completions.create(
                model=MODEL,
                messages=messages
            )
            result = chat_completion.choices[0].message.content
        except Exception as e:
            result = f"Ошибка: {str(e)}"

        if result:

            doc = Document()
            doc.add_heading('ShypagerAI – Диагноз', 0)
            doc.add_paragraph(result)

            temp_folder = os.path.join(os.getcwd(), "temp")
            if not os.path.exists(temp_folder):
                os.makedirs(temp_folder)

            username = session.get('username', 'user')
            safe_name = "".join([c if c.isalnum() else "_" for c in username])
            filename = f"diagnosis_{safe_name}.docx"
            filepath = os.path.join(temp_folder, filename)

            try:
                doc.save(filepath)
                session['word_path'] = filepath
                print(f"[DEBUG] Saved file at: {filepath}")
            except Exception as e:
                print(f"[ERROR] Failed to save docx: {e}")
                result += "\n[Ошибка при сохранении файла]"

    return render_template('osnov.html', result=result, username=session.get('username'), data_points=data_points)

@app.route('/download')
def download():
    filepath = session.get('word_path')
    print(f"[DEBUG] Downloading file: {filepath}")
    if filepath and os.path.exists(filepath):
        return send_file(
            filepath,
            as_attachment=True,
            download_name="diagnosis.docx",
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
    return "Нет файла для загрузки"

if __name__ == '__main__':
    app.run(debug=True) 